from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".json", ".yaml", ".yml", ".docx", ".pdf"}
MAX_FILE_BYTES = 5 * 1024 * 1024
START_MARKER = "PLAYBOOK_DATA_START"
END_MARKER = "PLAYBOOK_DATA_END"


def _document_text(data: bytes, extension: str) -> str:
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if extension == ".docx":
        document = Document(io.BytesIO(data))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(lines)
    return data.decode("utf-8")


def _structured_section(text: str) -> str:
    if START_MARKER in text and END_MARKER in text:
        return text.split(START_MARKER, 1)[1].split(END_MARKER, 1)[0].strip()
    stripped = text.strip()
    if stripped.startswith("{") and stripped.endswith("}"):
        return stripped
    raise ValueError(f"PDF and DOCX playbooks must contain a structured section between {START_MARKER} and {END_MARKER}")


def _validate_playbook(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ValueError("Playbook content must be an object")
    if not isinstance(payload.get("version"), (str, int, float)):
        raise ValueError("Playbook version is required")
    levels = payload.get("levels")
    if not isinstance(levels, dict):
        raise ValueError("Playbook levels must be an object")
    required_levels = {"monitor", "investigate", "escalate"}
    if set(levels) != required_levels:
        raise ValueError("Playbook levels must contain exactly monitor, investigate and escalate")
    required_action_fields = {"category", "action", "owner", "timeframe", "requires_approval"}
    for level, actions in levels.items():
        if not isinstance(actions, list) or not actions:
            raise ValueError(f"The {level} level must contain at least one action")
        for action in actions:
            if not isinstance(action, dict) or not required_action_fields.issubset(action):
                raise ValueError(f"Each {level} action must contain {', '.join(sorted(required_action_fields))}")
            if not isinstance(action["requires_approval"], bool):
                raise ValueError("requires_approval must be true or false")
    guidance = payload.get("condition_guidance")
    if not isinstance(guidance, dict) or not all(isinstance(items, list) for items in guidance.values()):
        raise ValueError("condition_guidance must map condition codes to lists")
    return payload


def parse_playbook_document(filename: str, data: bytes) -> dict[str, Any]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Supported playbook formats are PDF, DOCX, YAML and JSON")
    if not data:
        raise ValueError("The uploaded playbook is empty")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Playbook files must be 5 MB or smaller")
    text = _document_text(data, extension)
    structured = text if extension in {".json", ".yaml", ".yml"} else _structured_section(text)
    try:
        payload = json.loads(structured) if structured.lstrip().startswith("{") else yaml.safe_load(structured)
    except (json.JSONDecodeError, yaml.YAMLError) as error:
        raise ValueError("The structured playbook content is not valid JSON or YAML") from error
    playbook = _validate_playbook(payload)
    return {
        "valid": True,
        "source": {
            "filename": Path(filename).name,
            "format": extension.lstrip("."),
            "size_bytes": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
        },
        "summary": {
            "version": str(playbook["version"]),
            "conditions": sorted(playbook["condition_guidance"]),
            "action_counts": {level: len(actions) for level, actions in playbook["levels"].items()},
        },
        "playbook": playbook,
    }
