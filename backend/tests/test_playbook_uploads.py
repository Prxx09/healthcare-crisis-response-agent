import asyncio
import io
import json
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

from docx import Document
from app.config import Settings
from app.playbook_uploads import extract_text, extract_actions, structure_section, ingest, query_document, index_actions
from app.action_engine import action_rows
from datetime import datetime, timezone


def settings(**kwargs):
    return Settings(supabase_url='https://example.supabase.co', supabase_publishable_key='public',
                    pinecone_api_key='test', _env_file=None, **kwargs)


class UploadTests(unittest.IsolatedAsyncioTestCase):
    async def test_existing_package_without_llm(self):
        path = Path(__file__).resolve().parents[2] / 'playbooks/response_playbooks.json'
        data = path.read_bytes()
        actions, method = await extract_actions(path.name, data, extract_text(path.name, data), settings())
        self.assertEqual(method, 'structured')
        self.assertEqual({a['condition_code'] for a in actions}, {'ili', 'age', 'dengue'})
        self.assertTrue(all(a['requires_approval'] for a in actions))

    async def test_exact_dengue_section_file_has_nine_actions_without_ai(self):
        from app.playbook_documents import parse_playbook_document
        data = (Path(__file__).parent / 'fixtures/dengue_sections.json').read_bytes()
        validation = parse_playbook_document('dengue.json', data)
        self.assertEqual(validation['summary']['version'],'1.0')
        self.assertEqual(validation['summary']['action_counts'],{'monitor':3,'investigate':3,'escalate':3})
        with patch('app.playbook_uploads.structure_section',new=AsyncMock(side_effect=AssertionError('AI must not be called'))):
            actions, method = await extract_actions('dengue.json',data,data.decode(),settings())
        self.assertEqual(method,'structured'); self.assertEqual(len(actions),9)
        self.assertEqual({a['condition_code'] for a in actions},{'dengue'})
        self.assertEqual({a['section'] for a in actions},{1,2,3})
        monitor = next(a for a in actions if a['action'].startswith('For routine Dengue monitoring, the Surveillance'))
        self.assertEqual(monitor['owner'],'Surveillance lead')
        self.assertEqual(monitor['timeframe'],'within 24 hours')
        self.assertEqual(monitor['review_notes'],[])

    async def test_invalid_recognized_section_format_never_falls_back_to_ai(self):
        data=json.dumps({'document_version':'1','condition_code':'dengue','sections':[{'heading':'Unknown level','instructions':['Review dengue counts.']}]}).encode()
        with patch('app.playbook_uploads.structure_section',new=AsyncMock()) as model:
            with self.assertRaisesRegex(ValueError,'Section 1 must declare'):
                await extract_actions('guide.json',data,data.decode(),settings())
        model.assert_not_called()

    async def test_missing_schema_version_keeps_validation_error(self):
        data=json.dumps({'levels':{}}).encode()
        with self.assertRaisesRegex(ValueError,'Playbook version is required'):
            await extract_actions('guide.json',data,data.decode(),settings())

    async def test_index_failure_preserves_extraction_and_stage(self):
        data=(Path(__file__).parent / 'fixtures/dengue_sections.json').read_bytes()
        repo=MagicMock(); repo.privileged_get=AsyncMock(return_value=[])
        repo.post=AsyncMock(return_value=[{'id':'abc'}]);repo.patch=AsyncMock(return_value=[])
        with patch('app.playbook_uploads.upload_source',new=AsyncMock()), patch('app.playbook_uploads.index_actions',side_effect=RuntimeError('Index dimension mismatch')):
            with self.assertRaisesRegex(RuntimeError,'Pinecone indexing failed: Index dimension mismatch'):
                await ingest(repo,settings(),'dengue.json',data)
        updates=[call.args[2] for call in repo.patch.call_args_list]
        self.assertEqual(updates[0]['extraction_method'],'structured')
        self.assertEqual(len(updates[0]['actions']),9)
        self.assertEqual(updates[-1]['status'],'failed')
        self.assertIn('Pinecone indexing',updates[-1]['error_message'])

    async def test_http_error_identifies_storage_stage_without_leaking_headers(self):
        import httpx
        request=httpx.Request('POST','https://example.invalid',headers={'Authorization':'secret'})
        error=httpx.HTTPStatusError('sensitive detail',request=request,response=httpx.Response(403,request=request))
        repo=MagicMock();repo.privileged_get=AsyncMock(return_value=[])
        repo.post=AsyncMock(return_value=[{'id':'abc'}]);repo.patch=AsyncMock(return_value=[])
        with patch('app.playbook_uploads.upload_source',new=AsyncMock(side_effect=error)):
            with self.assertRaisesRegex(RuntimeError,'Source storage failed: Upstream service returned HTTP 403') as result:
                await ingest(repo,settings(),'guide.json',b'{"guidance":"Dengue operational guidance"}')
        self.assertNotIn('secret',str(result.exception))
        self.assertNotIn('sensitive',str(result.exception))

    def test_docx_tables_and_paragraphs(self):
        doc = Document(); doc.add_paragraph('Dengue investigation response playbook')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Review pending dengue laboratory results.'
        data = io.BytesIO(); doc.save(data)
        self.assertIn('pending dengue', extract_text('guide.docx', data.getvalue()))

    def test_empty_and_unsupported_rejected(self):
        for filename, data in [('file.exe', b'x'*40), ('file.json', b'')]:
            with self.assertRaises(ValueError): extract_text(filename, data)

    async def test_unstructured_needs_ai(self):
        data = json.dumps({'guidance':'Dengue: Review pending tests for the affected hospital.'}).encode()
        with self.assertRaises(RuntimeError):
            await extract_actions('guide.json', data, data.decode(), settings())

    async def model_result(self, quote):
        response = MagicMock(); response.json.return_value = {'choices':[{'message':{'content':json.dumps({'actions':[{
            'condition_code':'dengue','alert_level':'investigate','evidence_quote':quote,
            'owner':'Invented owner','timeframe':'within 2 hours','category':'investigation'}]})}}]}
        client = AsyncMock(); client.post.return_value = response
        context = MagicMock(); context.__aenter__ = AsyncMock(return_value=client); context.__aexit__ = AsyncMock()
        with patch('app.playbook_uploads.httpx.AsyncClient', return_value=context):
            return await structure_section('Dengue: Review pending dengue laboratory results.', 1, settings(ai_provider='groq',groq_api_key='test',groq_model='test'))

    async def test_hallucinated_quote_rejected(self):
        with self.assertRaises(ValueError): await self.model_result('Administer an invented treatment immediately.')

    async def test_owner_and_deadline_not_in_source_are_flagged(self):
        actions = await self.model_result('Review pending dengue laboratory results.')
        self.assertEqual(actions[0]['owner'], 'Incident Commander')
        self.assertEqual(actions[0]['timeframe'], 'After review')
        self.assertEqual(len(actions[0]['review_notes']), 2)

    async def test_duplicate_does_not_reindex(self):
        repo = MagicMock(); repo.privileged_get = AsyncMock(return_value=[{'status':'indexed','id':'abc'}])
        with patch('app.playbook_uploads.index_actions') as index:
            result = await ingest(repo, settings(), 'guide.json', b'Dengue operational response action instructions.')
        self.assertTrue(result['duplicate']); index.assert_not_called()

    async def test_failure_saved_and_not_activated(self):
        repo = MagicMock(); repo.privileged_get = AsyncMock(return_value=[])
        repo.post = AsyncMock(return_value=[{'id':'abc'}]); repo.patch = AsyncMock(return_value=[])
        with patch('app.playbook_uploads.upload_source', new=AsyncMock()), patch('app.playbook_uploads.extract_actions', new=AsyncMock(side_effect=ValueError('No supported action'))):
            with self.assertRaisesRegex(ValueError, 'Parsing and action extraction failed: No supported action'):
                await ingest(repo, settings(), 'guide.json', b'Dengue operational response action instructions.')
        self.assertEqual(repo.patch.call_args.args[2]['status'], 'failed')
        self.assertFalse(repo.post.call_args.args[1]['active'])

    def test_retrieval_uses_canonical_action_and_scope(self):
        action = {'action_id':'1','condition_code':'dengue','alert_level':'investigate','action':'Source instruction'}
        client = MagicMock(); client.inference.embed.return_value = {'data':[{'values':[0.1]}]}
        index = MagicMock(); index.query.return_value = {'matches':[{'id':'1','score':0.9,'metadata':{'text':'INJECTED'}},{'id':'999','score':1}]}
        with patch('app.playbook_uploads._pinecone_client',return_value=client), patch('app.playbook_uploads._index',return_value=index):
            result = query_document(settings(), {'id':'doc','filename':'guide.pdf','index_namespace':'isolated','actions':[action]}, 'dengue','investigate','query')
        self.assertEqual(len(result),1); self.assertEqual(result[0]['action'],'Source instruction')
        self.assertEqual(index.query.call_args.kwargs['namespace'],'isolated')

    async def test_uploaded_plan_replaces_bundled_actions_and_falls_back(self):
        from app.pinecone_rag import select_response_plan_with_rag
        action = {'action_id':'0','source_document_id':'a'*64,'action':'Verify pending dengue tests.'}
        with patch('app.pinecone_rag.retrieve_playbook_context', new=AsyncMock(return_value={'enabled':True,'status':'ready','matches':[]})), patch('app.playbook_uploads.uploaded_response_plan', new=AsyncMock(return_value={'action_source':'uploaded','actions':[action]})):
            result = await select_response_plan_with_rag(settings(pinecone_rag_enabled=True), 'dengue','investigate','pending labs')
        self.assertEqual(result['action_source'],'uploaded')
        self.assertEqual(result['actions'],[action])
        with patch('app.pinecone_rag.retrieve_playbook_context', new=AsyncMock(return_value={'enabled':True,'status':'ready','matches':[]})), patch('app.playbook_uploads.uploaded_response_plan', new=AsyncMock(return_value=None)):
            result = await select_response_plan_with_rag(settings(pinecone_rag_enabled=True), 'dengue','investigate')
        self.assertEqual(result['action_source'],'bundled')
        self.assertTrue(result['actions'])

    async def test_only_active_indexed_documents_are_searched(self):
        from app.playbook_uploads import uploaded_actions
        repo = MagicMock(); repo.privileged_get=AsyncMock(return_value=[])
        with patch('app.playbook_uploads.SupabaseRepository',return_value=repo):
            self.assertEqual(await uploaded_actions(settings(),'ili','monitor','query'),[])
        self.assertEqual(repo.privileged_get.call_args.args[0],'playbook_condition_rules')
        self.assertEqual(repo.privileged_get.call_args.args[1]['condition_code'],'eq.ili')

    def test_first_upload_skips_cleanup_and_creates_vectors(self):
        client=MagicMock();client.inference.embed.return_value={'data':[{'values':[0.1]}]}
        index=MagicMock();index.describe_index_stats.return_value={'namespaces':{},'dimension':1}
        a={'action_id':'0','condition_code':'dengue','alert_level':'monitor','action':'Check reports','owner':'Surveillance lead','timeframe':'within 24 hours'}
        with patch('app.playbook_uploads._pinecone_client',return_value=client),patch('app.playbook_uploads._index',return_value=index):
            self.assertEqual(index_actions(settings(),{'id':'doc','index_namespace':'new','actions':[a]}),1)
        index.delete.assert_not_called()
        index.upsert.assert_called_once()
        self.assertEqual(index.upsert.call_args.kwargs['namespace'],'new')

    def test_retry_only_clears_the_existing_document_namespace(self):
        client=MagicMock();index=MagicMock()
        index.describe_index_stats.return_value={'namespaces':{'existing':{'vector_count':9},'other':{'vector_count':5}}}
        with patch('app.playbook_uploads._pinecone_client',return_value=client),patch('app.playbook_uploads._index',return_value=index):
            index_actions(settings(),{'id':'doc','index_namespace':'existing','actions':[]})
        index.delete.assert_called_once_with(delete_all=True,namespace='existing')

    def test_missing_index_is_not_silently_ignored(self):
        client=MagicMock();index=MagicMock()
        error=Exception('missing index');error.status=404
        index.describe_index_stats.side_effect=error
        with patch('app.playbook_uploads._pinecone_client',return_value=client),patch('app.playbook_uploads._index',return_value=index):
            with self.assertRaisesRegex(RuntimeError,r'Index/namespace lookup failed: Exception \(HTTP 404\)'):
                index_actions(settings(),{'id':'doc','index_namespace':'new','actions':[]})
        index.upsert.assert_not_called()

    def test_batch_embedding_and_source_lineage(self):
        a = {'action_id':'1','condition_code':'dengue','alert_level':'investigate','action':'Verify tests','owner':'Lab lead','timeframe':'After review','category':'investigation'}
        client=MagicMock();client.inference.embed.return_value={'data':[{'values':[0.1]}]};index=MagicMock()
        with patch('app.playbook_uploads._pinecone_client',return_value=client),patch('app.playbook_uploads._index',return_value=index):
            self.assertEqual(index_actions(settings(),{'id':'doc','index_namespace':'isolated','actions':[a]}),1)
        self.assertEqual(index.upsert.call_args.kwargs['vectors'][0]['id'],'1')
        row = action_rows(1,[{**a,'source_document_id':'doc'}],datetime.now(timezone.utc))[0]
        self.assertEqual(row['source_document_id'],'doc')


if __name__ == '__main__': unittest.main()
